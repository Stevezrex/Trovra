#!/usr/bin/env python3
"""Builds the Week 7 marking workbook from the merged CSV (merge.mjs output).

Usage: python3 build_workbook.py trials_merged.csv Marking_Workbook.xlsx [--fake]

- Learner messages are copied exactly as exported. The workbook never edits them.
- The instructor chooses which submitted trial to mark (Learners sheet). Nothing is selected automatically.
- Rubric scores are typed into the yellow cells. All totals are formulas.
"""
import csv, math, re, sys
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

src, dst = sys.argv[1], sys.argv[2]
FAKE = '--fake' in sys.argv

rows = list(csv.DictReader(open(src, encoding='utf-8-sig', newline='')))
if not rows:
    sys.exit('The CSV has no rows.')

# ---------- styles ----------
FONT = 'Arial'
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_title = Font(name=FONT, size=14, bold=True, color='0E1729')
f_small = Font(name=FONT, size=9, color='4A546B')
f_input = Font(name=FONT, size=10, color='0000FF')
f_red = Font(name=FONT, size=10, bold=True, color='A63A2B')
NAVY = PatternFill('solid', fgColor='0E1729')
SLATE = PatternFill('solid', fgColor='16233D')
AMBER = PatternFill('solid', fgColor='F2A33A')
YELLOW = PatternFill('solid', fgColor='FFFF00')
PAPER = PatternFill('solid', fgColor='F4F2EC')
GREEN = PatternFill('solid', fgColor='D6EADF')
GREY = PatternFill('solid', fgColor='EBE8DF')
thin = Side(style='thin', color='D9D5C8')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

def style_range(ws, ref, font=None, fill=None, align=None, border=None):
    for row in ws[ref]:
        for c in (row if isinstance(row, tuple) else (row,)):
            if font: c.font = font
            if fill: c.fill = fill
            if align: c.alignment = align
            if border: c.border = border

def est_height(text, width_chars, line_pt=12.5, cap=400):
    lines = 0
    for ln in str(text).split('\n'):
        lines += max(1, math.ceil(len(ln) / max(1, width_chars)))
    return min(cap, max(15, lines * line_pt + 4))

# ---------- rubric text (from the plan, section 10 and 11) ----------
COMMON = [
    'Paraphrases the facts accurately',
    'Acknowledges the specific impact without arguing about emotion',
    'Does not match hostility',
    'Separates verified facts from assumptions',
    'Asks relevant questions when needed, or proceeds without unnecessary questions when the facts are sufficient',
    'Explains realistic options and limits',
    'Confirms the next action, owner and timing',
]
CASE_CRIT = {
    'A': ['Checks the order and tracking before promising', 'Promises no unconfirmed delivery date', 'Gives an update time separate from the completion time'],
    'B': ['Recognizes the earlier poor contact', 'Treats it as a defect, not a change of mind', 'States the remedy and timeline in writing'],
    'C': ['Keeps investigating the duplicate charge despite the insults', 'After the first insult, states a boundary naming the behaviour and a condition', 'After the relapse, pauses and transfers to the supervisor, keeps the Payments step and time, and does not retaliate'],
    'D': ['Clarifies the loss and the amount asked', 'Explains authority honestly without refusing to investigate', 'Escalates with a usable summary, owner and time, and does not suggest approval'],
}
NOTE_CRIT = [
    'Issue and wanted outcome are stated as facts',
    'Verified and unconfirmed facts are separated',
    'Recorded action matches the transcript',
    'Next step, owner and time match what the customer was told',
    'Conduct or escalation entry matches the case requirements',
]
EXPECT = {
    'A': 'Acceptance condition: the learner has checked the order or tracking, has stated a trace with an update time, has explained the refund route that applies if no date can be confirmed, and has promised no delivery date.',
    'B': 'Acceptance condition: the learner treats it as a defect and not a change of mind, asks for the photo or short video if it is not yet sent, states replacement or refund with the policy timeline, and has not applied the change of mind rule.',
    'C': 'Acceptance condition: the customer does not accept before the relapse line. After it, the learner has verified the two debits and given a Payments next step with owner and time. If the learner paused the chat and transferred to the shift supervisor, the case closes as transferred.',
    'D': 'Acceptance condition: the learner clarifies the loss and the amount asked, explains the limit honestly, arranges the correction, escalates to the Claims reviewer with a usable summary, and gives a review time. The learner does not suggest approval is likely.',
}
CASE_NAMES = {'A': 'Case A: delayed delivery (chat)', 'B': 'Case B: defective item (email)', 'C': 'Case C: duplicate charge with insults (chat)', 'D': 'Case D: compensation above authority (email)'}
NOTE_LABELS = ['Issue', 'Key facts', 'Verification', 'Action', 'Outcome', 'Next step', 'Owner and date', 'Conduct or escalation']

# ---------- organise data ----------
trials = OrderedDict()
for r in rows:
    key = (r['learner_email'], int(r['trial_no']))
    trials.setdefault(key, {'meta': r, 'cases': {}})['cases'][r['case_id']] = r
learners = OrderedDict()
for (email, no), t in trials.items():
    learners.setdefault(email, {'name': t['meta']['learner_name'], 'trials': []})['trials'].append((no, t))
submitted = [(k, t) for k, t in trials.items() if t['meta']['trial_status'] == 'submitted']

wb = Workbook()

# =====================================================================
# Guide
# =====================================================================
g = wb.active
g.title = 'Guide'
g.sheet_view.showGridLines = False
g.column_dimensions['A'].width = 3
g.column_dimensions['B'].width = 30
g.column_dimensions['C'].width = 100
g['B2'] = 'Week 7 marking workbook: Trovra support desk simulation'; g['B2'].font = f_title
r = 3
if FAKE:
    g['B3'] = 'FAKE DATA. Every record in this file was generated by a script. Use it to test the workflow, not to mark learners.'
    g['B3'].font = f_red
    r = 4
g.cell(row=r + 1, column=2, value='Certificate in Customer Service and Virtual Assistance. Week 7: difficult customers and conflict resolution. Checkpoint: role-play rubric.').font = f_base
steps = [
    ('1. Choose the trial', 'On the Learners sheet, type 1, 2 or 3 in the yellow Selected trial cell for each learner. Nothing is chosen for you: the workbook does not pick the latest or the highest-scoring trial. Only submitted trials can be marked.'),
    ('2. Read the evidence', 'Case sheets list every submitted trial. Rows for the selected trial turn green and show SELECTED. Filter column D on SELECTED to see only those. Read the Transcript and Note columns. The Transcripts sheet shows the same transcripts one message per row.'),
    ('3. Score', 'Type scores into the yellow cells. Reply criteria R1 to R10: 0 (not shown), 1 (partly shown), 2 (clearly shown). Note criteria N1 to N5: 0 or 1. Totals calculate by themselves.'),
    ('4. Check totals', 'The Learners sheet adds the four case totals into a mark out of 100 for the selected trial. The Trials sheet shows a total for every trial.'),
    ('Engine evidence', 'Flags, Signals and Heat history come from the rule engine. They point you to things worth checking, such as an overpromise, a boundary, or a missed window. They do not award marks and they can miss valid phrasings. Score from the transcript.'),
    ('Transcripts are untouched', 'Learner messages are copied exactly as stored. Do not edit them. Times are West Africa Time (WAT).'),
    ('Calibration', 'Before marking the cohort, two people (or you twice, a day apart) score the same two sample transcripts. Record the totals on the Calibration sheet and discuss any criterion where scores differ by more than 1 mark.'),
    ('Marks', 'Each case: 20 marks for replies (10 criteria x 2) plus 5 marks for the service recovery note. Four cases: 100 marks per trial.'),
    ('Cell colours', 'Yellow fill with blue text: you type here. Green row: selected trial. Everything else is locked in your head only: the sheets are not protected, so avoid typing over them.'),
]
rr = r + 3
for k, v in steps:
    g.cell(row=rr, column=2, value=k).font = f_bold
    c = g.cell(row=rr, column=3, value=v); c.font = f_base; c.alignment = wrap_top
    g.cell(row=rr, column=2).alignment = wrap_top
    g.row_dimensions[rr].height = est_height(v, 105, 13, 120)
    rr += 1
g.cell(row=rr + 1, column=2, value='Source file').font = f_bold
g.cell(row=rr + 1, column=3, value=f'{src}: {len(trials)} trial(s), {len(submitted)} submitted, {len(learners)} learner(s).').font = f_base

# =====================================================================
# Rubric
# =====================================================================
ru = wb.create_sheet('Rubric')
ru.sheet_view.showGridLines = False
ru.column_dimensions['A'].width = 8
ru.column_dimensions['B'].width = 70
for col in 'CDEF':
    ru.column_dimensions[col].width = 40
ru['A1'] = 'Rubric'; ru['A1'].font = f_title
ru['A3'] = 'Reply criteria: 10 per case, scored 0, 1 or 2 (20 marks)'; ru['A3'].font = f_bold
hdr = ['No.', 'Common criteria (all cases)']
for i, h in enumerate(hdr):
    c = ru.cell(row=4, column=1 + i, value=h); c.font = f_head; c.fill = NAVY
for i, t in enumerate(COMMON):
    ru.cell(row=5 + i, column=1, value=i + 1).font = f_base
    c = ru.cell(row=5 + i, column=2, value=t); c.font = f_base; c.alignment = wrap_top
ru['A13'] = 'Case criteria (8 to 10)'; ru['A13'].font = f_bold
for i, h in enumerate(['No.', 'Case A', 'Case B', 'Case C', 'Case D']):
    c = ru.cell(row=14, column=1 + i, value=h); c.font = f_head; c.fill = NAVY
for j in range(3):
    ru.cell(row=15 + j, column=1, value=8 + j).font = f_base
    for ci, cid in enumerate('ABCD'):
        c = ru.cell(row=15 + j, column=2 + ci, value=CASE_CRIT[cid][j]); c.font = f_base; c.alignment = wrap_top
    ru.row_dimensions[15 + j].height = 66
ru['A19'] = 'Service recovery note: 5 criteria, 1 mark each (5 marks)'; ru['A19'].font = f_bold
for i, h in enumerate(['No.', 'Criterion']):
    c = ru.cell(row=20, column=1 + i, value=h); c.font = f_head; c.fill = NAVY
for i, t in enumerate(NOTE_CRIT):
    ru.cell(row=21 + i, column=1, value=i + 1).font = f_base
    ru.cell(row=21 + i, column=2, value=t).font = f_base
ru['A27'] = 'Acceptance conditions (what a full resolution needs)'; ru['A27'].font = f_bold
for i, cid in enumerate('ABCD'):
    ru.cell(row=28 + i, column=1, value=cid).font = f_bold
    c = ru.cell(row=28 + i, column=2, value=EXPECT[cid].replace('Acceptance condition: ', '')); c.font = f_base; c.alignment = wrap_top
    ru.merge_cells(start_row=28 + i, start_column=2, end_row=28 + i, end_column=5)
    ru.row_dimensions[28 + i].height = 44

# =====================================================================
# Case sheets (one row per submitted trial)
# =====================================================================
# column layout
IDENT = ['Learner', 'Email', 'Trial', 'Selected?']
EVID = ['Outcome', 'Missed windows', 'Learner replies', 'Transcript', 'Service recovery note', 'Engine flags (evidence only)', 'Signals by reply (evidence only)', 'Heat history (evidence only)', 'Case facts']
WIDTH = {'Learner': 20, 'Email': 28, 'Trial': 7, 'Selected?': 12, 'Outcome': 13, 'Missed windows': 9, 'Learner replies': 9, 'Transcript': 78, 'Service recovery note': 52, 'Engine flags (evidence only)': 34, 'Signals by reply (evidence only)': 60, 'Heat history (evidence only)': 30, 'Case facts': 34}
FIRST = 6  # first data row
case_cols = {}  # per case: dict of column indexes
last_row = {}

def note_text(r):
    parts = []
    for n, lab in enumerate(NOTE_LABELS, start=1):
        parts.append(f'{n}. {lab}: {r.get(f"note_{n}", "")}')
    status = r['note_status'] + (' (time expired, saved text)' if r['note_time_expired'] == 'yes' else '')
    return f'[{status}]\n' + '\n'.join(parts)

n_data = max(1, len(submitted))
for cid in 'ABCD':
    ws = wb.create_sheet(f'Case {cid}')
    ws.sheet_view.showGridLines = False
    ws['A1'] = CASE_NAMES[cid]; ws['A1'].font = f_title
    ws['A2'] = EXPECT[cid]; ws['A2'].font = f_small
    ws['A3'] = 'Yellow cells: type 0, 1 or 2 for R1 to R10, and 0 or 1 for N1 to N5. Rows for the trial you selected on the Learners sheet turn green.'; ws['A3'].font = f_small
    cols = IDENT + EVID
    c0 = len(cols) + 1
    r_cols = list(range(c0, c0 + 10))
    r_tot = c0 + 10
    n_cols = list(range(r_tot + 1, r_tot + 6))
    n_tot = r_tot + 6
    c_tot = n_tot + 1
    filled = c_tot + 1
    case_cols[cid] = dict(email=2, trial=3, sel=4, rtot=r_tot, ntot=n_tot, ctot=c_tot, filled=filled)
    # headers (row 5)
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=5, column=i, value=h); c.font = f_head; c.fill = NAVY; c.alignment = center; c.border = box
        ws.column_dimensions[L(i)].width = WIDTH[h]
    crit = COMMON + CASE_CRIT[cid]
    for k, ci in enumerate(r_cols):
        c = ws.cell(row=5, column=ci, value=f'R{k + 1}. {crit[k]}'); c.font = f_head; c.fill = SLATE; c.alignment = center; c.border = box
        ws.column_dimensions[L(ci)].width = 17
    c = ws.cell(row=5, column=r_tot, value='Reply total (of 20)'); c.font = f_bold; c.fill = AMBER; c.alignment = center; c.border = box
    ws.column_dimensions[L(r_tot)].width = 11
    for k, ci in enumerate(n_cols):
        c = ws.cell(row=5, column=ci, value=f'N{k + 1}. {NOTE_CRIT[k]}'); c.font = f_head; c.fill = SLATE; c.alignment = center; c.border = box
        ws.column_dimensions[L(ci)].width = 17
    for ci, label in ((n_tot, 'Note total (of 5)'), (c_tot, 'Case total (of 25)'), (filled, 'Entries filled (of 15)')):
        c = ws.cell(row=5, column=ci, value=label); c.font = f_bold; c.fill = AMBER; c.alignment = center; c.border = box
        ws.column_dimensions[L(ci)].width = 12
    ws.row_dimensions[5].height = 118

    # data rows
    first, last = FIRST, FIRST + n_data - 1
    r = FIRST
    for (email, no), t in submitted:
        m = t['meta']; cr = t['cases'][cid]
        vals = [m['learner_name'], email, no, None,
                cr['case_outcome'], int(cr['missed_windows']), int(cr['learner_replies']),
                cr['transcript'], note_text(cr), cr['flags'], cr['signals'], cr['heat_history'], cr['case_facts']]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v); c.font = f_base; c.alignment = wrap_top; c.border = box
        ws.cell(row=r, column=4, value=f'=IF(IFERROR(INDEX(Learners!$D$6:$D$500,MATCH($B{r},Learners!$B$6:$B$500,0)),"")=$C{r},"SELECTED","")').font = f_bold
        ws.cell(row=r, column=4).alignment = Alignment(vertical='top', horizontal='center')
        for ci in r_cols + n_cols:
            c = ws.cell(row=r, column=ci); c.fill = YELLOW; c.font = f_input; c.border = box; c.alignment = Alignment(horizontal='center', vertical='top')
        rng_r = f'{L(r_cols[0])}{r}:{L(r_cols[-1])}{r}'
        rng_n = f'{L(n_cols[0])}{r}:{L(n_cols[-1])}{r}'
        ws.cell(row=r, column=r_tot, value=f'=IF(COUNT({rng_r})=0,"",SUM({rng_r}))')
        ws.cell(row=r, column=n_tot, value=f'=IF(COUNT({rng_n})=0,"",SUM({rng_n}))')
        ws.cell(row=r, column=c_tot, value=f'=IF(COUNT({rng_r},{rng_n})=0,"",SUM({rng_r},{rng_n}))')
        ws.cell(row=r, column=filled, value=f'=COUNT({rng_r},{rng_n})')
        for ci in (r_tot, n_tot, c_tot, filled):
            c = ws.cell(row=r, column=ci); c.font = f_bold; c.alignment = Alignment(horizontal='center', vertical='top'); c.border = box
        h = max(est_height(cr['transcript'], 88), est_height(note_text(cr), 58), est_height(cr['signals'], 66), 60)
        ws.row_dimensions[r].height = h
        r += 1
    last_row[cid] = last
    # validation
    dv2 = DataValidation(type='whole', operator='between', formula1='0', formula2='2', allow_blank=True, showErrorMessage=True,
                         errorTitle='Reply score', error='Enter 0, 1 or 2.')
    dv1 = DataValidation(type='whole', operator='between', formula1='0', formula2='1', allow_blank=True, showErrorMessage=True,
                         errorTitle='Note score', error='Enter 0 or 1.')
    ws.add_data_validation(dv2); ws.add_data_validation(dv1)
    dv2.add(f'{L(r_cols[0])}{first}:{L(r_cols[-1])}{last}')
    dv1.add(f'{L(n_cols[0])}{first}:{L(n_cols[-1])}{last}')
    ws.conditional_formatting.add(f'A{first}:{L(c0 - 1)}{last}', FormulaRule(formula=[f'$D{first}="SELECTED"'], fill=GREEN))
    ws.freeze_panes = f'E{FIRST}'
    ws.auto_filter.ref = f'A5:{L(filled)}{last}'
    ws.sheet_properties.tabColor = 'F2A33A'

def rng(cid, col, abs_=True):
    last = last_row[cid]
    return f"'Case {cid}'!${L(case_cols[cid][col])}${FIRST}:${L(case_cols[cid][col])}${last}"

# =====================================================================
# Learners
# =====================================================================
ls = wb.create_sheet('Learners', 1)
ls.sheet_view.showGridLines = False
ls['A1'] = 'Learners and marks'; ls['A1'].font = f_title
ls['A2'] = 'Choose the trial to mark in the yellow column. Marks come from the case sheets for that trial. A blank selection means that learner is not yet chosen.'; ls['A2'].font = f_small
heads = ['Learner', 'Email', 'Trials submitted', 'Selected trial', 'Selection check',
         'A replies', 'A note', 'B replies', 'B note', 'C replies', 'C note', 'D replies', 'D note',
         'Total (of 100)', 'Entries filled (of 60)']
widths = [22, 30, 14, 10, 24, 10, 8, 10, 8, 10, 8, 10, 8, 14, 14]
for i, (h, w) in enumerate(zip(heads, widths), start=1):
    c = ls.cell(row=5, column=i, value=h); c.font = f_head; c.fill = NAVY; c.alignment = center; c.border = box
    ls.column_dimensions[L(i)].width = w
ls.row_dimensions[5].height = 34
r = 6
for email, info in learners.items():
    subs = sorted(no for no, t in info['trials'] if t['meta']['trial_status'] == 'submitted')
    ls.cell(row=r, column=1, value=info['name'])
    ls.cell(row=r, column=2, value=email)
    ls.cell(row=r, column=3, value=', '.join(map(str, subs)) if subs else 'none')
    sel = ls.cell(row=r, column=4); sel.fill = YELLOW; sel.font = f_input; sel.alignment = Alignment(horizontal='center')
    ls.cell(row=r, column=5, value=(
        f'=IF($D{r}="","Not selected",IF(COUNTIFS(Trials!$B$6:$B$500,$B{r},Trials!$C$6:$C$500,$D{r},Trials!$E$6:$E$500,"submitted")=1,"OK","Not a submitted trial"))'))
    col = 6
    for cid in 'ABCD':
        ls.cell(row=r, column=col, value=f'=IF($D{r}="","",SUMIFS({rng(cid, "rtot")},{rng(cid, "email")},$B{r},{rng(cid, "trial")},$D{r}))')
        ls.cell(row=r, column=col + 1, value=f'=IF($D{r}="","",SUMIFS({rng(cid, "ntot")},{rng(cid, "email")},$B{r},{rng(cid, "trial")},$D{r}))')
        col += 2
    ls.cell(row=r, column=14, value=f'=IF($D{r}="","",SUM(F{r}:M{r}))')
    parts = '+'.join(f'SUMIFS({rng(cid, "filled")},{rng(cid, "email")},$B{r},{rng(cid, "trial")},$D{r})' for cid in 'ABCD')
    ls.cell(row=r, column=15, value=f'=IF($D{r}="","",{parts})')
    for i in range(1, 16):
        c = ls.cell(row=r, column=i)
        if i != 4: c.font = f_bold if i == 14 else f_base
        c.border = box
        if i >= 6: c.alignment = Alignment(horizontal='center')
    r += 1
l_last = r - 1
dvs = DataValidation(type='list', formula1='"1,2,3"', allow_blank=True, showErrorMessage=True, errorTitle='Selected trial', error='Choose 1, 2 or 3.')
ls.add_data_validation(dvs); dvs.add(f'D6:D{l_last}')
ls.conditional_formatting.add(f'E6:E{l_last}', FormulaRule(formula=['$E6="Not a submitted trial"'], font=Font(name=FONT, bold=True, color='A63A2B')))
ls.freeze_panes = 'C6'
ls.auto_filter.ref = f'A5:O{l_last}'
ls.sheet_properties.tabColor = '0E1729'
# Learners range must cover the lookup ranges used in the case sheets ($6:$500). Values beyond l_last are blank.

# =====================================================================
# Trials
# =====================================================================
tr = wb.create_sheet('Trials', 2)
tr.sheet_view.showGridLines = False
tr['A1'] = 'All trials'; tr['A1'].font = f_title
tr['A2'] = 'Every trial that exists, submitted or open. Open trials cannot be marked. Totals show the marks entered so far for each trial.'; tr['A2'].font = f_small
heads = ['Learner', 'Email', 'Trial', 'Fact set', 'Status', 'Shift started (WAT)', 'Submitted (WAT)', 'Resets', 'Offline gaps', 'Selected for marking?', 'Total entered (of 100)']
widths = [22, 30, 7, 9, 11, 26, 26, 8, 44, 14, 14]
for i, (h, w) in enumerate(zip(heads, widths), start=1):
    c = tr.cell(row=5, column=i, value=h); c.font = f_head; c.fill = NAVY; c.alignment = center; c.border = box
    tr.column_dimensions[L(i)].width = w
tr.row_dimensions[5].height = 34
r = 6
for (email, no), t in trials.items():
    m = t['meta']
    vals = [m['learner_name'], email, no, int(m['fact_set']), m['trial_status'], m['shift_started_wat'], m['submitted_wat'], int(m['reset_count']), m['offline_gaps'] or '']
    for i, v in enumerate(vals, start=1):
        c = tr.cell(row=r, column=i, value=v); c.font = f_base; c.border = box; c.alignment = wrap_top
    tr.cell(row=r, column=10, value=f'=IF(IFERROR(INDEX(Learners!$D$6:$D$500,MATCH($B{r},Learners!$B$6:$B$500,0)),"")=$C{r},"SELECTED","")')
    if m['trial_status'] == 'submitted':
        parts = '+'.join(f'SUMIFS({rng(cid, "ctot")},{rng(cid, "email")},$B{r},{rng(cid, "trial")},$C{r})' for cid in 'ABCD')
        tr.cell(row=r, column=11, value=f'={parts}')
    else:
        tr.cell(row=r, column=11, value='Open: not markable')
    for i in (10, 11):
        c = tr.cell(row=r, column=i); c.font = f_bold; c.border = box; c.alignment = Alignment(horizontal='center', vertical='top')
    if m['offline_gaps']:
        tr.row_dimensions[r].height = est_height(m['offline_gaps'], 50, 13, 120)
    r += 1
t_last = r - 1
tr.conditional_formatting.add(f'A6:K{t_last}', FormulaRule(formula=['$J6="SELECTED"'], fill=GREEN))
tr.conditional_formatting.add(f'A6:K{t_last}', FormulaRule(formula=['$E6="open"'], fill=GREY))
tr.freeze_panes = 'C6'
tr.auto_filter.ref = f'A5:K{t_last}'
# Trials!B/C/E are used by the Learners selection check

# =====================================================================
# Transcripts (one message per row)
# =====================================================================
tx = wb.create_sheet('Transcripts')
tx.sheet_view.showGridLines = False
tx['A1'] = 'Transcripts, one message per row'; tx['A1'].font = f_title
tx['A2'] = 'Copied exactly as exported. Filter by learner, trial and case.'; tx['A2'].font = f_small
heads = ['Learner', 'Email', 'Trial', 'Case', 'Time (WAT)', 'Speaker', 'Note', 'Message']
widths = [20, 28, 7, 7, 11, 11, 26, 100]
for i, (h, w) in enumerate(zip(heads, widths), start=1):
    c = tx.cell(row=4, column=i, value=h); c.font = f_head; c.fill = NAVY; c.alignment = center
    tx.column_dimensions[L(i)].width = w
START = re.compile(r'^\[(\d{1,2}:\d{2}[ap]m)\] (CUSTOMER|AGENT)([^\n]*?):\n?', re.M)
r = 5
for (email, no), t in submitted:
    for cid in 'ABCD':
        cr = t['cases'][cid]
        text = cr['transcript']
        ms = list(START.finditer(text))
        for i, m in enumerate(ms):
            end = ms[i + 1].start() if i + 1 < len(ms) else len(text)
            body = text[m.end():end].rstrip('\n')
            subj = ''
            sm = re.match(r'   Subject: ([^\n]*)\n', body)
            if sm:
                subj = sm.group(1); body = body[sm.end():]
            note = (m.group(3).strip() + (' | ' if m.group(3).strip() and subj else '') + (f'Subject: {subj}' if subj else '')).strip()
            vals = [t['meta']['learner_name'], email, no, cid, m.group(1), m.group(2).title(), note, body]
            for ci, v in enumerate(vals, start=1):
                c = tx.cell(row=r, column=ci, value=v); c.font = f_base; c.alignment = wrap_top
            tx.row_dimensions[r].height = est_height(body, 108, 13, 300)
            r += 1
tx.freeze_panes = 'A5'
tx.auto_filter.ref = f'A4:H{max(5, r - 1)}'

# =====================================================================
# Calibration
# =====================================================================
cb = wb.create_sheet('Calibration')
cb.sheet_view.showGridLines = False
cb['A1'] = 'Calibration'; cb['A1'].font = f_title
cb['A2'] = 'Score two sample transcripts before marking the cohort. Enter each marker\'s total out of 100 (or score the same marker twice). Review any difference above 5 marks.'; cb['A2'].font = f_small
heads = ['Sample', 'Learner and trial scored', 'Marker 1 total', 'Marker 2 total', 'Difference', 'Criteria to discuss']
widths = [10, 36, 14, 14, 12, 60]
for i, (h, w) in enumerate(zip(heads, widths), start=1):
    c = cb.cell(row=4, column=i, value=h); c.font = f_head; c.fill = NAVY; c.alignment = center; c.border = box
    cb.column_dimensions[L(i)].width = w
for rr_ in (5, 6):
    cb.cell(row=rr_, column=1, value=rr_ - 4).font = f_base
    for ci in (2, 3, 4, 6):
        c = cb.cell(row=rr_, column=ci); c.fill = YELLOW; c.font = f_input; c.border = box
    d = cb.cell(row=rr_, column=5, value=f'=IF(OR(C{rr_}="",D{rr_}=""),"",ABS(C{rr_}-D{rr_}))'); d.font = f_bold; d.border = box
cb['A8'] = 'Example of the format: Sample 1 | Chidi Eze, trial 2 | 78 | 74 | 4 | R6 on Case C'; cb['A8'].font = f_small

# order sheets
order = ['Guide', 'Learners', 'Trials', 'Case A', 'Case B', 'Case C', 'Case D', 'Transcripts', 'Rubric', 'Calibration']
wb._sheets = [wb[n] for n in order]
wb.save(dst)
print(f'Saved {dst}: {len(learners)} learners, {len(trials)} trials, {len(submitted)} submitted')
